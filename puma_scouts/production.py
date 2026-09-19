from __future__ import annotations
import asyncio,csv,json,os,re,statistics,time,xml.etree.ElementTree as ET
from collections import Counter
from pathlib import Path
from typing import Any,Callable
from openpyxl import Workbook,load_workbook
from .classic_report import save_classic
from .identity_map import load_discovery_sources,remember_discovery_sources
from .models import IdentityConfidence,ProductMission,Verdict,ScanHealth,ScanReport
from .price_score import assess_price_market
from .scouts.base import repair_discovery_scope

def _text(v:Any)->str:return "" if v is None else str(v).strip()
def _currency_code(v:Any)->str:
    text=_text(v).upper().replace(".","")
    aliases={"":"UAH","UAH":"UAH","ГРН":"UAH","₴":"UAH","HUA":"UAH","USD":"USD","$":"USD","US$":"USD","EUR":"EUR","€":"EUR"}
    return aliases.get(text,text or "UAH")
def _is_uah_currency(v:Any)->bool:return _currency_code(v)=="UAH"
ALIASES={
    "article":{"артикул","article","sku","код","vendorcode","vendor_code"},
    "name":{"назва","название","name","товар","product","title"},
    "brand":{"бренд","brand","виробник","vendor"},
    "model":{"модель","model","mpn"},
    "own_price":{"ціна","цена","price","price_uah","вартість","стоимость"},
    "category":{"категорія","категория","category","categoryname","category_name"},
    "supplier":{"постачальник","поставщик","supplier","vendor_name"},
    "old_price":{"стара ціна","старая цена","old_price","oldprice"},
    "discount":{"знижка","знижка, %","скидка","скидка, %","discount","discount_pct"},
}
def _mission(data:dict[str,Any],fallback:str)->ProductMission|None:
    norm={str(k).strip().lower():_text(v) for k,v in data.items()}
    picked={k:next((norm[a] for a in names if norm.get(a)),"") for k,names in ALIASES.items()}
    if not picked["name"]:return None
    source={k:v for k,v in picked.items() if k!="article" and v}
    return ProductMission(article=picked["article"] or fallback,source_data=source)
def _read_xlsx(path:str,limit=None):
    wb=load_workbook(path,data_only=True,read_only=True);ws=wb.active;header=0;cols={}
    for r in range(1,min(ws.max_row,30)+1):
        vals={_text(ws.cell(r,c).value).lower():c for c in range(1,ws.max_column+1)}
        found={k:next((vals[n] for n in names if n in vals),0) for k,names in ALIASES.items()}
        if found["name"] or (found["article"] and sum(bool(x) for x in found.values())>=2):header,cols=r,found;break
    if not header:raise ValueError("Потрібна колонка Назва/Name; Артикул, Бренд і Модель можуть бути окремими колонками")
    out=[]
    for r in range(header+1,ws.max_row+1):
        data={k:_text(ws.cell(r,c).value) if c else "" for k,c in cols.items()};m=_mission(data,f"ROW-{r}")
        if m:out.append(m)
        if limit and len(out)>=limit:break
    return out
def _read_csv(path:str,limit=None):
    raw=Path(path).read_bytes();text=None
    for enc in ("utf-8-sig","utf-8","cp1251"):
        try:text=raw.decode(enc);break
        except UnicodeDecodeError:pass
    if text is None:raise ValueError("Не вдалося визначити кодування CSV")
    try:dialect=csv.Sniffer().sniff(text[:8192],delimiters=",;\t|")
    except csv.Error:dialect=csv.excel;dialect.delimiter=";"
    out=[]
    for i,row in enumerate(csv.DictReader(text.splitlines(),dialect=dialect),2):
        m=_mission(row,f"ROW-{i}")
        if m:out.append(m)
        if limit and len(out)>=limit:break
    return out
def _tag(el):return el.tag.rsplit('}',1)[-1].lower()
def _read_xml(path:str,limit=None):
    root=ET.parse(path).getroot();out=[]
    for i,e in enumerate([e for e in root.iter() if _tag(e) in {"offer","product","item"}],1):
        data={_tag(c):_text(c.text) for c in e.iter() if c is not e and _text(c.text)}
        for k,v in e.attrib.items():data.setdefault(k.lower(),v)
        m=_mission(data,f"ROW-{i}")
        if m:out.append(m)
        if limit and len(out)>=limit:break
    return out
def read_missions(path:str,limit:int|None=None)->list[ProductMission]:
    ext=Path(path).suffix.lower()
    if ext==".xlsx":out=_read_xlsx(path,limit)
    elif ext==".csv":out=_read_csv(path,limit)
    elif ext in {".yml",".xml"}:out=_read_xml(path,limit)
    else:raise ValueError("Підтримуються XLSX, CSV, YML та XML")
    if not out:raise ValueError("У файлі не знайдено товарів із назвою")
    return out

def scouts(selected:list[str]|None=None):
    from .scouts.catalog import HotlineScout,PromScout
    from .scouts.epicentr import EpicentrScout
    from .scouts.web_shops import WebShopsScout
    all_scouts={
        "epicentr":EpicentrScout(timeout=8,max_candidates_per_query=10),
        "prom":PromScout(timeout=8,max_candidates_per_query=10),
        "hotline":HotlineScout(timeout=8,max_candidates_per_query=10),
        "web_shops":WebShopsScout(timeout=8,max_candidates_per_query=12,max_per_domain=2),
    }
    wanted=set(selected or all_scouts);return [s for k,s in all_scouts.items() if k in wanted]

SOURCE_WALL_TIMEOUT=float(os.getenv("PUMA_SOURCE_WALL_TIMEOUT","70"))
SOURCE_HARD_TIMEOUT=max(
    SOURCE_WALL_TIMEOUT,
    float(os.getenv("PUMA_SOURCE_HARD_TIMEOUT",str(SOURCE_WALL_TIMEOUT*2))),
)
PRODUCT_CONCURRENCY=max(1,min(int(os.getenv("PUMA_PRODUCT_CONCURRENCY","4")),4))

def _env_flag(name:str,default:str="0")->bool:
    return os.getenv(name,default).strip().casefold() in {"1","true","yes","on"}

async def _scan_source(
    scout,
    mission:ProductMission,
    wall_timeout:float=SOURCE_WALL_TIMEOUT,
    hard_timeout:float|None=None,
):
    """Run one source with bounded recovery while preserving the legacy return shape."""
    soft=max(0.01,float(wall_timeout))
    hard=max(soft,float(SOURCE_HARD_TIMEOUT if hard_timeout is None else hard_timeout))
    task=asyncio.create_task(scout.scan(mission))
    try:
        try:
            report=await asyncio.wait_for(asyncio.shield(task),timeout=soft)
            return scout,report
        except asyncio.TimeoutError:
            remaining=max(0.01,hard-soft)
            try:
                report=await asyncio.wait_for(asyncio.shield(task),timeout=remaining)
                report.errors.append(f"soft_timeout_{soft:g}s_recovered")
                return scout,report
            except asyncio.TimeoutError:
                task.cancel()
                await asyncio.gather(task,return_exceptions=True)
                return scout,ScanReport(
                    article=mission.article,marketplace=scout.marketplace,health=ScanHealth.ACCESS_LIMITED,
                    errors=[f"hard_timeout_{hard:g}s"]
                )
    except Exception as exc:
        if not task.done():
            task.cancel()
            await asyncio.gather(task,return_exceptions=True)
        return scout,ScanReport(
            article=mission.article,marketplace=scout.marketplace,health=ScanHealth.SCOUT_ERROR,
            errors=[f"{type(exc).__name__}: {exc}"]
        )

async def _timed_scan_source(scout,mission:ProductMission):
    started=time.perf_counter()
    source,report=await _scan_source(scout,mission)
    return source,report,time.perf_counter()-started

def _title_signature(title:str)->str:
    t=title.casefold();t=re.sub(r"\([^)]*\)$","",t);t=re.sub(r"\b(?:монітор|монитор)\b"," ",t);return re.sub(r"[^a-zа-яіїєґ0-9]+"," ",t).strip()
def _prom_identity(x):return (x["article"],x["source"],x["price"],_title_signature(x["found_title"]))

def _reason_summary(report)->str:
    reasons=Counter()
    for item in report.offers:
        for reason in list(item.conflicts or [])+list(item.rejection_reasons or []):
            if reason: reasons[reason]+=1
    return " | ".join(f"{reason} ×{count}" for reason,count in reasons.most_common(5))

async def scan_detailed(mission:ProductMission,selected:list[str]|None=None,scout_pool:list[Any]|None=None,force_discovery:bool=False)->tuple[list[dict[str,Any]],dict[str,Any]]:
    if scout_pool is not None:
        wanted=set(selected or [])
        active_scouts=[
            s for s in scout_pool
            if not wanted or getattr(getattr(s,"marketplace",None),"value",None) in wanted
        ]
    else:
        active_scouts=scouts(selected)
    rows=[];diagnostics={}
    with repair_discovery_scope(force_discovery):
        results=await asyncio.gather(*[_timed_scan_source(s,mission) for s in active_scouts])
    for scout,report,elapsed in results:
        source=scout.marketplace.value
        accepted=0
        verdicts=Counter()
        identities=Counter()
        price_integrity=Counter()
        for item in report.offers:
            verdicts[item.verdict.value]+=1
            if item.identity_confidence:identities[item.identity_confidence.value]+=1
            if item.verdict!=Verdict.PASS or item.identity_confidence not in {None,IdentityConfidence.CONFIRMED,IdentityConfidence.PROBABLE}:continue
            o=item.offer
            currency=_currency_code(o.currency)
            if o.price is not None and currency!="UAH":
                price_integrity[f"non-UAH price excluded: {currency}"]+=1
                continue
            accepted+=1
            rows.append({
                "article":mission.article,"name":mission.source_data.get("name",""),"brand":mission.source_data.get("brand",""),
                "model":mission.source_data.get("model",""),"own_price":mission.source_data.get("own_price",""),"source":source,
                "price":float(o.price) if o.price is not None else None,"currency":currency,"availability":o.availability or "",
                "found_title":o.title,"url":str(o.url),"match":round(item.score,3),
                "identity_confidence":item.identity_confidence.value if item.identity_confidence else "LEGACY_PASS",
                "domain":o.attributes.get("source_domain","") or str(o.url).split('/')[2],"health":report.health.value
            })
        reason_summary=_reason_summary(report)
        if price_integrity:
            price_reasons=" | ".join(f"{reason} ×{count}" for reason,count in price_integrity.items())
            reason_summary=" | ".join(x for x in (reason_summary,price_reasons) if x)
        diagnostics[source]={
            "health":report.health.value,"queries":report.queries_generated,"pages":report.pages_scanned,
            "candidates":report.candidates_collected,"seen":report.candidates_seen,"accepted":accepted,
            "pass":verdicts.get("PASS",0),"conflict":verdicts.get("CONFLICT",0),"reject":verdicts.get("REJECT",0),
            "ambiguous":identities.get("AMBIGUOUS",0),"identity_conflict":identities.get("CONFLICT",0),
            "reasons":reason_summary,"errors":list(report.errors or []),"elapsed_seconds":round(elapsed,3),
            "metrics":dict(report.metrics or {}),
        }
    dedup={}
    for x in rows:
        if x["source"]=="web_shops":key=(x["article"],x["source"],x["domain"],x["price"])
        elif x["source"]=="prom":key=_prom_identity(x)
        else:key=(x["article"],x["source"],x["url"])
        current=dedup.get(key)
        if current is None or x["match"]>current["match"]:dedup[key]=x
    return list(dedup.values()),diagnostics

async def scan(mission:ProductMission,selected:list[str]|None=None)->list[dict[str,Any]]:
    pool=scouts(selected)
    try:
        rows,_=await scan_detailed(mission,selected,scout_pool=pool)
        return rows
    finally:
        await asyncio.gather(*(s.aclose() for s in pool if hasattr(s,"aclose")),return_exceptions=True)

def _market_summary(offers):
    groups={}
    for x in offers:
        if x.get("price") is None or not _is_uah_currency(x.get("currency","UAH")):continue
        groups.setdefault(x["source"],[]).append(x["price"])
    return " | ".join(f"{src}: "+", ".join(f"{price:g} грн"+(f" ×{prices.count(price)}" if prices.count(price)>1 else "") for price in sorted(set(prices))) for src,prices in sorted(groups.items()))

def product_report(missions:list[ProductMission],rows:list[dict[str,Any]],diagnostics_by_article:dict[str,dict[str,Any]]|None=None)->list[dict[str,Any]]:
    by_article={}
    for x in rows:by_article.setdefault(x["article"],[]).append(x)
    diagnostics_by_article=diagnostics_by_article or {};report=[]
    for m in missions:
        offers=by_article.get(m.article,[])
        conf=Counter(x.get("identity_confidence","LEGACY_PASS") for x in offers)
        market=assess_price_market(offers,m.source_data.get("own_price",""))
        report.append({
            "article":m.article,"name":m.source_data.get("name",""),"brand":m.source_data.get("brand",""),"model":m.source_data.get("model",""),
            "category":m.source_data.get("category",""),"supplier":m.source_data.get("supplier",""),"old_price":m.source_data.get("old_price",""),
            "discount":m.source_data.get("discount",""),"own_price":m.source_data.get("own_price",""),"status":"FOUND" if offers else "NOT_FOUND",
            "marketplaces":_market_summary(offers),"offers":len(offers),"sources":market["market_sources"],
            "min_price":market["market_min"],"median_price":market["market_median"],
            "avg_price":market["market_avg"],"max_price":market["market_max"],
            "confirmed":conf.get("CONFIRMED",0),"probable":conf.get("PROBABLE",0),
            "price_score":market["price_score"],"price_verdict":market["price_verdict"],
            "price_verdict_reason":market["price_verdict_reason"],"market_confidence":market["market_confidence"],
            "suspicious_price_count":market["suspicious_price_count"],"valid_offer_count":market["valid_offer_count"],"reserve_uah":market["reserve_uah"],
            "reserve_pct":market["reserve_pct"],"delta_median_pct":market["delta_median_pct"],
            "same_price_count":market["same_price_count"],"market_representatives":market["market_representatives"],
            "offer_rows":offers,"diagnostics":diagnostics_by_article.get(m.article,{})
        })
    return report

def save(rows:list[dict[str,Any]],output:str,missions:list[ProductMission]|None=None,diagnostics_by_article:dict[str,dict[str,Any]]|None=None)->None:
    missions=missions or [];products=product_report(missions,rows,diagnostics_by_article) if missions else [];wb=Workbook();overview=wb.active;overview.title="Результат"
    overview.append(["Артикул","Назва","Бренд","Модель","Вхідна ціна","Статус","Вердикт ціни","Причина","Price Score","Достовірність","Маркетплейс → ціна","Мін. ринку","Медіана","Середня","Макс. ринку","Карток","Джерел ринку","Підозрілих цін","Запас, грн","Запас, %","CONFIRMED","PROBABLE"])
    for p in products:overview.append([p[k] for k in ("article","name","brand","model","own_price","status","price_verdict","price_verdict_reason","price_score","market_confidence","marketplaces","min_price","median_price","avg_price","max_price","offers","sources","suspicious_price_count","reserve_uah","reserve_pct","confirmed","probable")])
    overview.freeze_panes="A2";overview.auto_filter.ref=overview.dimensions
    ws=wb.create_sheet("Пропозиції");keys=("article","name","brand","model","own_price","source","price","currency","price_status","price_reason","availability","found_title","url","match","identity_confidence","domain","health")
    ws.append(["Артикул","Назва","Бренд","Модель","Вхідна ціна","Джерело","Ціна","Валюта","Статус ціни","Причина ціни","Наявність","Знайдена назва","URL","Match","Identity Confidence","Домен","Health"]);defaults={"identity_confidence":"LEGACY_PASS","domain":"","health":"","own_price":"","price_status":"","price_reason":""}
    for x in rows:ws.append([x.get(k,defaults.get(k,"")) for k in keys])
    ws.freeze_panes="A2";ws.auto_filter.ref=ws.dimensions
    summary=wb.create_sheet("Ціна → картки");summary.append(["Артикул","Назва","Джерело","Ціна","Карток"]);groups=Counter((x["article"],x["name"],x["source"],x["price"]) for x in rows if x.get("price") is not None and _is_uah_currency(x.get("currency","UAH")))
    for key,count in sorted(groups.items(),key=lambda z:(z[0][0],z[0][2],z[0][3])):summary.append([*key,count])
    summary.freeze_panes="A2";summary.auto_filter.ref=summary.dimensions
    diag=wb.create_sheet("Discovery діагностика")
    diag.append(["Артикул","Товар","Джерело","Health","Секунд","Запитів","Сторінок","Кандидатів","Seen","Прийнято","PASS","CONFLICT","REJECT","AMBIGUOUS","ID CONFLICT","Identity URLs","Identity Hit","Identity Saved","Serper Attempts","Serper API","Serper Cache","Serper Rescued","Serper Success Query","Причини","Помилки"])
    name_by_article={m.article:m.source_data.get("name","") for m in missions}
    for article,sources in (diagnostics_by_article or {}).items():
        for src,d in sources.items():
            diag.append([
                article,name_by_article.get(article,""),src,d.get("health",""),d.get("elapsed_seconds",0),d.get("queries",0),d.get("pages",0),
                d.get("candidates",0),d.get("seen",0),d.get("accepted",0),d.get("pass",0),d.get("conflict",0),
                d.get("reject",0),d.get("ambiguous",0),d.get("identity_conflict",0),
                (d.get("metrics") or {}).get("identity_urls_loaded",0),
                bool((d.get("metrics") or {}).get("identity_refresh_hit",False)),
                (d.get("metrics") or {}).get("identity_urls_saved",0),
                (d.get("metrics") or {}).get("serper_queries_attempted",0),
                (d.get("metrics") or {}).get("serper_api_requests",0),
                (d.get("metrics") or {}).get("serper_cache_hits",0),
                bool((d.get("metrics") or {}).get("serper_rescued",False)),
                (d.get("metrics") or {}).get("serper_success_query",0),
                d.get("reasons","")," | ".join(d.get("errors") or [])
            ])
    diag.freeze_panes="A2";diag.auto_filter.ref=diag.dimensions;wb.save(output)


def _load_checkpoint(path:str|None,token:str|None,missions:list[ProductMission])->dict[int,dict[str,Any]]:
    if not path:return {}
    p=Path(path)
    if not p.is_file():return {}
    try:
        data=json.loads(p.read_text(encoding="utf-8"))
    except Exception:
        return {}
    articles=[m.article for m in missions]
    if data.get("version")!=1:return {}
    if str(data.get("checkpoint_token") or "")!=str(token or ""):return {}
    if list(data.get("articles") or [])!=articles:return {}
    completed={}
    for key,value in (data.get("completed") or {}).items():
        try:index=int(key)
        except (TypeError,ValueError):continue
        if not isinstance(value,dict) or index<0 or index>=len(missions):continue
        if value.get("article")!=missions[index].article:continue
        completed[index]=value
    return completed

def _save_checkpoint(path:str|None,token:str|None,missions:list[ProductMission],completed:dict[int,dict[str,Any]])->None:
    if not path:return
    p=Path(path);p.parent.mkdir(parents=True,exist_ok=True)
    payload={
        "version":1,
        "checkpoint_token":str(token or ""),
        "articles":[m.article for m in missions],
        "completed_indexes":sorted(completed),
        "completed":{str(i):completed[i] for i in sorted(completed)},
    }
    tmp=p.with_suffix(p.suffix+".tmp")
    tmp.write_text(json.dumps(payload,ensure_ascii=False,separators=(",",":")),encoding="utf-8")
    os.replace(tmp,p)

async def run(input_path:str,output_path:str,limit:int|None=None,selected:list[str]|None=None,progress_cb:Callable|None=None,classic_output_path:str|None=None,supplier:str="",checkpoint_path:str|None=None,checkpoint_token:str|None=None)->dict[str,Any]:
    missions=read_missions(input_path,limit)
    completed_state=_load_checkpoint(checkpoint_path,checkpoint_token,missions)
    scout_pool=scouts(selected)
    semaphore=asyncio.Semaphore(PRODUCT_CONCURRENCY)
    lock=asyncio.Lock()
    refresh_mode=_env_flag("PUMA_REFRESH_ONLY")
    repair_enabled=refresh_mode and _env_flag("PUMA_REFRESH_REPAIR_ENABLED")

    async def one(index:int,m:ProductMission):
        async with semaphore:
            async with lock:
                already_done=len(completed_state)
            if progress_cb:progress_cb(already_done,len(missions),m.source_data.get("name",m.article),f"Шукаємо реальні пропозиції… (паралельність {PRODUCT_CONCURRENCY})")
            found,diag=await scan_detailed(m,selected,scout_pool=scout_pool)
            if not refresh_mode:
                remember_discovery_sources(m,{x.get("source","") for x in found if x.get("source")})
            elif repair_enabled and not found:
                prior_sources=load_discovery_sources(m)
                if prior_sources:
                    allowed=set(selected or prior_sources)
                    repair_targets=[src for src in prior_sources if src in allowed]
                    if repair_targets:
                        repair_found,repair_diag=await scan_detailed(
                            m,repair_targets,scout_pool=scout_pool,force_discovery=True
                        )
                        repair_sources={x.get("source") for x in repair_found if x.get("source")}
                        for src in repair_targets:
                            before=diag.get(src,{})
                            after=repair_diag.get(src,before)
                            before_metrics=dict(before.get("metrics") or {})
                            after_metrics=dict(after.get("metrics") or {})
                            after_metrics["product_repair_attempted"]=True
                            after_metrics["product_repair_rescued"]=src in repair_sources
                            after_metrics["refresh_identity_urls_loaded_before_repair"]=int(
                                before_metrics.get("identity_urls_loaded") or 0
                            )
                            after_metrics["refresh_identity_urls_attempted_before_repair"]=int(
                                before_metrics.get("identity_urls_attempted") or 0
                            )
                            after_metrics["refresh_repair_required_before_repair"]=bool(
                                before_metrics.get("repair_required")
                            )
                            after_metrics["refresh_discovery_gap_before_repair"]=bool(
                                before_metrics.get("discovery_gap")
                            )
                            after["metrics"]=after_metrics
                            after["errors"]=list(before.get("errors") or [])+list(after.get("errors") or [])
                            after["elapsed_seconds"]=round(
                                float(before.get("elapsed_seconds") or 0)
                                +float(after.get("elapsed_seconds") or 0),3
                            )
                            diag[src]=after
                        if repair_found:
                            found=repair_found
            async with lock:
                completed_state[index]={"article":m.article,"found":found,"diagnostics":diag}
                _save_checkpoint(checkpoint_path,checkpoint_token,missions,completed_state)
                current=len(completed_state)
                if progress_cb:progress_cb(current,len(missions),m.source_data.get("name",m.article),"Товар перевірено")
            return index

    pending=[(i,m) for i,m in enumerate(missions) if i not in completed_state]
    if progress_cb and completed_state:
        progress_cb(len(completed_state),len(missions),"Відновлення",f"Продовжуємо з {len(completed_state)} з {len(missions)} вже перевірених товарів")
    if pending:
        await asyncio.gather(*[one(i,m) for i,m in pending])

    rows=[]
    diagnostics_by_article={}
    for i,m in enumerate(missions):
        item=completed_state.get(i) or {"found":[],"diagnostics":{}}
        rows.extend(item.get("found") or [])
        diagnostics_by_article[m.article]=item.get("diagnostics") or {}

    products=product_report(missions,rows,diagnostics_by_article)
    save(rows,output_path,missions,diagnostics_by_article)
    if classic_output_path:save_classic(products,classic_output_path,supplier=supplier)
    _save_checkpoint(checkpoint_path,checkpoint_token,missions,completed_state)
    if progress_cb:progress_cb(len(missions),len(missions),"Готово","Формуємо звіти…")
    found=sum(1 for p in products if p["status"]=="FOUND")
    confidence=Counter(x.get("identity_confidence","LEGACY_PASS") for x in rows)
    timing_samples={}
    for source_map in diagnostics_by_article.values():
        for src,d in source_map.items():
            timing_samples.setdefault(src,[]).append(float(d.get("elapsed_seconds") or 0))
    source_timing={
        src:{
            "samples":len(vals),
            "avg_seconds":round(sum(vals)/max(1,len(vals)),3),
            "max_seconds":round(max(vals),3) if vals else 0,
        }
        for src,vals in timing_samples.items()
    }
    source_discovery={}
    for source_map in diagnostics_by_article.values():
        for src,d in source_map.items():
            bucket=source_discovery.setdefault(src,{
                "products":0,"identity_refresh_hits":0,"identity_urls_loaded":0,"identity_urls_attempted":0,
                "identity_urls_saved":0,"repair_required_products":0,"discovery_gap_products":0,
                "repair_attempted_products":0,"repair_rescued_products":0,
                "serper_queries_attempted":0,"serper_api_requests":0,"serper_cache_hits":0,
                "serper_rescued_products":0,"serper_first_query_success":0,"serper_second_query_success":0,
            })
            bucket["products"]+=1
            metrics=d.get("metrics") or {}
            bucket["identity_refresh_hits"]+=int(bool(metrics.get("identity_refresh_hit")))
            bucket["identity_urls_loaded"]+=int(metrics.get("identity_urls_loaded") or 0)
            bucket["identity_urls_attempted"]+=int(metrics.get("identity_urls_attempted") or 0)
            bucket["identity_urls_saved"]+=int(metrics.get("identity_urls_saved") or 0)
            bucket["repair_required_products"]+=int(bool(metrics.get("repair_required")))
            bucket["discovery_gap_products"]+=int(bool(metrics.get("discovery_gap")))
            bucket["repair_attempted_products"]+=int(bool(metrics.get("product_repair_attempted")))
            bucket["repair_rescued_products"]+=int(bool(metrics.get("product_repair_rescued")))
            bucket["serper_queries_attempted"]+=int(metrics.get("serper_queries_attempted") or 0)
            bucket["serper_api_requests"]+=int(metrics.get("serper_api_requests") or 0)
            bucket["serper_cache_hits"]+=int(metrics.get("serper_cache_hits") or 0)
            bucket["serper_rescued_products"]+=int(bool(metrics.get("serper_rescued")))
            success_query=int(metrics.get("serper_success_query") or 0)
            if success_query==1:bucket["serper_first_query_success"]+=1
            elif success_query==2:bucket["serper_second_query_success"]+=1

    planned_sources=[getattr(getattr(s,"marketplace",None),"value",None) for s in scout_pool]
    planned_sources=[src for src in planned_sources if src]
    if not planned_sources:
        planned_sources=sorted({x.get("source") for x in rows if x.get("source")})
    source_sets={}
    for m in missions:source_sets[m.article]=set()
    for x in rows:
        if x.get("price") is None or not _is_uah_currency(x.get("currency","UAH")):continue
        source_sets.setdefault(x["article"],set()).add(x["source"])
    coverage_by_source={
        src:{
            "products_with_price":sum(1 for values in source_sets.values() if src in values),
            "pct":round(sum(1 for values in source_sets.values() if src in values)/max(1,len(missions))*100,1),
        }
        for src in planned_sources
    }
    source_count_distribution={
        str(n):sum(1 for values in source_sets.values() if len(values)==n)
        for n in range(0,len(planned_sources)+1)
    }
    all_sources_products=sum(
        1 for values in source_sets.values()
        if all(src in values for src in planned_sources)
    )
    serper_clients=[getattr(s,"serper",None) for s in scout_pool if getattr(s,"serper",None) is not None]
    serper_api_requests=sum(int(getattr(s,"api_requests",0) or 0) for s in serper_clients)
    serper_cache_hits=sum(int(getattr(s,"cache_hits",0) or 0) for s in serper_clients)
    serper_configured=bool(os.getenv("SERPER_API_KEY","").strip())
    serper_enabled=bool(serper_clients) and all(bool(getattr(s,"enabled",False)) for s in serper_clients)
    serper_disabled_reasons=sorted({str(getattr(s,"disabled_reason","") or "") for s in serper_clients if getattr(s,"disabled_reason","")})
    price_verdicts=Counter(p.get("price_verdict","") for p in products if p.get("price_verdict"))
    repair_products_attempted=sum(
        1 for source_map in diagnostics_by_article.values()
        if any(bool((d.get("metrics") or {}).get("product_repair_attempted")) for d in source_map.values())
    )
    repair_products_rescued=sum(
        1 for source_map in diagnostics_by_article.values()
        if any(bool((d.get("metrics") or {}).get("product_repair_rescued")) for d in source_map.values())
    )
    web_products=[{k:v for k,v in p.items() if k!="offer_rows"}|{"offers_detail":p["offer_rows"]} for p in products]
    result={
        "products":len(missions),"offers":len(rows),"found_products":found,
        "found_pct":round(found/max(1,len(missions))*100,1),"sources":len({x["source"] for x in rows}),
        "confirmed":confidence.get("CONFIRMED",0),"probable":confidence.get("PROBABLE",0),
        "product_concurrency":PRODUCT_CONCURRENCY,"source_timing":source_timing,
        "refresh_only":refresh_mode,
        "repair_enabled":repair_enabled,
        "repair_products_attempted":repair_products_attempted,
        "repair_products_rescued":repair_products_rescued,
        "serper_api_requests":serper_api_requests,"serper_cache_hits":serper_cache_hits,
        "serper_configured":serper_configured,"serper_enabled":serper_enabled,
        "serper_disabled_reasons":serper_disabled_reasons,
        "source_discovery":source_discovery,
        "coverage_by_source":coverage_by_source,
        "all_sources_products":all_sources_products,
        "all_sources_pct":round(all_sources_products/max(1,len(missions))*100,1),
        "source_count_distribution":source_count_distribution,
        "price_verdicts":dict(price_verdicts),
        "product_results":web_products,
    }
    await asyncio.gather(*(s.aclose() for s in scout_pool if hasattr(s,"aclose")),return_exceptions=True)
    return result

def run_sync(input_path:str,output_path:str,**kwargs):return asyncio.run(run(input_path,output_path,**kwargs))
